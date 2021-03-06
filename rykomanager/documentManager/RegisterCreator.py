from rykomanager.documentManager.DocumentCreator import DocumentCreator
from rykomanager.documentManager.DatabaseManager import DatabaseManager
from docx import Document
from os import path
from datetime import datetime
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from rykomanager.DateConverter import DateConverter
from rykomanager import config_parser


class RegisterCreator(DocumentCreator):
    CELL_TO_MERGE_MARK = "MERGE"

    def __init__(self, program_id):
        self.date = datetime.today().strftime('%d-%m-%Y')
        self.program = DatabaseManager.get_program(program_id)
        self.contracts = DatabaseManager.get_contracts(program_id)
        self.program_semester = self.program.get_current_semester()
        self.year = self.program.school_year
        self.records_to_merge = []
        output_directory = config_parser.get('Directories', 'current_program')
        DocumentCreator.__init__(self, config_parser.get('DocTemplates', 'register'), output_directory)

    def create(self):
        self.generate("Rejestr_{}.docx".format(self.date))

    def _prepare_school_data(self):
        no = 1
        for contract in self.contracts:
            self.fill_data_in_table(contract, no, False)

            # Fill additional fields with annex information if needed any
            for annex in DatabaseManager.get_annex(contract.program_id, contract.contract_no):
                self.fill_data_in_table(annex, no, True)
            no += 1

    def fill_data_in_table(self, contract, no, is_annex):
        record_dict = dict()
        if not is_annex:
            record_dict['no'] = str(no)
            record_dict['contract_info'] = "{}/{}".format(contract.contract_no, contract.contract_year)
            record_dict['school_name'] = contract.school.name
            record_dict['school_nip'] = contract.school.nip
            record_dict['school_address'] = contract.school.address
            record_dict['school_city'] = contract.school.city
            record_dict['school_regon'] = contract.school.regon
            record_dict['school_phone'] = contract.school.phone
            record_dict['school_email'] = contract.school.email
        else:
            record_dict['no'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_name'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_nip'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_address'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_city'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_regon'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_phone'] = RegisterCreator.CELL_TO_MERGE_MARK
            record_dict['school_email'] = RegisterCreator.CELL_TO_MERGE_MARK
            validity_date = DateConverter.to_date(contract.validity_date)
            record_dict['contract_info'] = "{}/{}".format(contract.contract_no.split("_")[0], contract.contract_year)
            record_dict['annex_info'] = "{}*".format(DateConverter.to_string(validity_date, "%d-%m-%Y"))

        record_dict['kids_milk'] = str(contract.dairy_products)
        record_dict['kids_fruitveg'] = str(contract.fruitVeg_products)
        self.records_to_merge.append(record_dict)

    def generate(self, new_doc_name, gen_pdf=True):
        self._prepare_school_data()
        self.document.merge_rows('no', self.records_to_merge)
        self.document.merge(
            semester_no=self.program_semester,
            school_year=self.year,
            date=self.date
        )
        DocumentCreator.generate(self, new_doc_name, False)
        generated_file = path.join(self.output_directory, new_doc_name)
        RegisterCreator._merge_cells(generated_file, RegisterCreator.CELL_TO_MERGE_MARK)
        DocumentCreator.generate_pdf(generated_file, self.output_directory)

    @staticmethod
    def _merge_cells(file_with_table_to_merge, mark):
        document = Document(file_with_table_to_merge)
        for table in document.tables:
            for col in range(0, len(table.columns)):
                for row in range(0, len(table.rows)):
                    if table.cell(row, col).text == mark:
                        table.cell(row, col).text = ""
                        merged = table.cell(row - 1, col).merge(table.cell(row, col))
                        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.allow_autofit = True
        document.save(file_with_table_to_merge)
